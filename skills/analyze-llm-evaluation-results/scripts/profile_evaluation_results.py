#!/usr/bin/env python3
"""Read-only profiler for LLM evaluation result tables.

The caller must confirm field meanings and pass/fail values before using this
script. It never edits the input and never sends data over the network.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable


class UserError(Exception):
    pass


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def folded(value: Any) -> str:
    return text(value).casefold()


def is_missing(value: Any, missing_values: set[str]) -> bool:
    return value is None or text(value) == "" or folded(value) in missing_values


def read_csv(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    raw = path.read_bytes()
    decoded = None
    used_encoding = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = raw.decode(encoding)
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        raise UserError("无法按 UTF-8 或 GB18030 解码；请显式转换编码后重试。")
    try:
        dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    reader = csv.DictReader(decoded.splitlines(), dialect=dialect)
    if not reader.fieldnames:
        raise UserError("文件没有表头。")
    rows = [dict(row) for row in reader]
    return rows, {"format": "csv", "encoding": used_encoding, "delimiter": dialect.delimiter}


def records_from_json(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        records = value
    elif isinstance(value, dict):
        records = None
        for key in ("data", "results", "items", "records"):
            if isinstance(value.get(key), list):
                records = value[key]
                break
        if records is None:
            records = [value]
    else:
        raise UserError("JSON 顶层必须是对象、对象数组，或包含 data/results/items/records 数组。")
    if not all(isinstance(row, dict) for row in records):
        raise UserError("记录数组中存在非对象元素。")
    return [dict(row) for row in records]


def read_json(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise UserError(f"JSON 读取失败：{exc}") from exc
    return records_from_json(value), {"format": "json"}


def read_jsonl(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise UserError(f"JSONL 第 {number} 行无效：{exc}") from exc
        if not isinstance(value, dict):
            raise UserError(f"JSONL 第 {number} 行不是对象。")
        rows.append(dict(value))
    return rows, {"format": "jsonl"}


def read_xlsx(path: Path, sheet: str | None) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise UserError(
            "读取 .xlsx 需要本地安装 openpyxl；请先获得用户许可再安装，"
            "或将指定工作表另存为 CSV。"
        ) from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet:
        if sheet not in workbook.sheetnames:
            raise UserError(f"找不到工作表 {sheet!r}；可选：{workbook.sheetnames}")
        worksheet = workbook[sheet]
    else:
        worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return [], {"format": "xlsx", "sheet": worksheet.title}
    headers = [text(value) or f"unnamed_{index + 1}" for index, value in enumerate(raw_headers)]
    if len(headers) != len(set(headers)):
        raise UserError("Excel 表头存在重复列名，请先确认列口径。")
    rows = [dict(zip(headers, values)) for values in iterator]
    return rows, {"format": "xlsx", "sheet": worksheet.title}


def read_rows(path: Path, sheet: str | None) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        return read_csv(path)
    if suffix == ".json":
        return read_json(path)
    if suffix in (".jsonl", ".ndjson"):
        return read_jsonl(path)
    if suffix == ".xlsx":
        return read_xlsx(path, sheet)
    if suffix == ".xls":
        raise UserError("旧版 .xls 不受脚本支持；请在本地另存为 .xlsx 或 CSV。")
    raise UserError(f"不支持的格式：{suffix}")


def ordered_columns(rows: list[dict[str, Any]]) -> list[str]:
    seen: set[str] = set()
    columns: list[str] = []
    for row in rows:
        for key in row:
            name = str(key)
            if name not in seen:
                seen.add(name)
                columns.append(name)
    return columns


def as_number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = float(text(value))
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def rounded(value: float | None) -> float | None:
    return None if value is None else round(value, 6)


def numeric_profile(rows: Iterable[dict[str, Any]], column: str, missing: set[str]) -> dict[str, Any]:
    values: list[float] = []
    invalid = 0
    for row in rows:
        raw = row.get(column)
        if is_missing(raw, missing):
            continue
        number = as_number(raw)
        if number is None:
            invalid += 1
        else:
            values.append(number)
    result: dict[str, Any] = {"valid_n": len(values), "invalid_n": invalid}
    if not values:
        return result
    values.sort()
    result.update(
        {
            "mean": rounded(statistics.fmean(values)),
            "median": rounded(statistics.median(values)),
            "min": rounded(values[0]),
            "max": rounded(values[-1]),
            "stdev": rounded(statistics.stdev(values)) if len(values) > 1 else None,
        }
    )
    if len(values) > 1:
        q1, _, q3 = statistics.quantiles(values, n=4, method="inclusive")
        result.update({"q1": rounded(q1), "q3": rounded(q3)})
    return result


def wilson(successes: int, total: int, z: float = 1.96) -> list[float] | None:
    if total == 0:
        return None
    p = successes / total
    denominator = 1 + z * z / total
    centre = (p + z * z / (2 * total)) / denominator
    margin = z * math.sqrt((p * (1 - p) + z * z / (4 * total)) / total) / denominator
    return [round(max(0.0, centre - margin), 6), round(min(1.0, centre + margin), 6)]


def make_classifier(pass_values: set[str], fail_values: set[str], missing: set[str]):
    def classify(value: Any) -> bool | None:
        if is_missing(value, missing):
            return None
        key = folded(value)
        if key in pass_values:
            return True
        if key in fail_values:
            return False
        return None

    return classify


def pass_profile(
    rows: Iterable[dict[str, Any]], column: str, classify, missing: set[str]
) -> dict[str, Any]:
    passed = failed = missing_n = unknown = 0
    unknown_values: Counter[str] = Counter()
    for row in rows:
        value = row.get(column)
        if is_missing(value, missing):
            missing_n += 1
            continue
        outcome = classify(value)
        if outcome is True:
            passed += 1
        elif outcome is False:
            failed += 1
        else:
            unknown += 1
            unknown_values[text(value)] += 1
    total = passed + failed
    return {
        "eligible_n": total,
        "passed_n": passed,
        "failed_n": failed,
        "missing_n": missing_n,
        "unrecognized_n": unknown,
        "unrecognized_values": dict(unknown_values.most_common(20)),
        "pass_rate": rounded(passed / total) if total else None,
        "wilson_95": wilson(passed, total),
    }


def group_rows(rows: list[dict[str, Any]], column: str, missing: set[str]):
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        value = row.get(column)
        label = "<MISSING>" if is_missing(value, missing) else text(value)
        groups[label].append(row)
    return groups


def nominal_kappa(pairs: list[tuple[str, str]]) -> float | None:
    if not pairs:
        return None
    total = len(pairs)
    observed = sum(left == right for left, right in pairs) / total
    left_counts = Counter(left for left, _ in pairs)
    right_counts = Counter(right for _, right in pairs)
    expected = sum(left_counts[key] * right_counts[key] for key in set(left_counts) | set(right_counts))
    expected /= total * total
    if math.isclose(expected, 1.0):
        return 1.0 if math.isclose(observed, 1.0) else None
    return rounded((observed - expected) / (1 - expected))


def pearson(pairs: list[tuple[float, float]]) -> float | None:
    if len(pairs) < 2:
        return None
    xs = [item[0] for item in pairs]
    ys = [item[1] for item in pairs]
    mean_x, mean_y = statistics.fmean(xs), statistics.fmean(ys)
    numerator = sum((x - mean_x) * (y - mean_y) for x, y in pairs)
    denominator = math.sqrt(sum((x - mean_x) ** 2 for x in xs) * sum((y - mean_y) ** 2 for y in ys))
    return rounded(numerator / denominator) if denominator else None


def agreement_profile(
    rows: list[dict[str, Any]], human: str, judge: str, missing: set[str]
) -> dict[str, Any]:
    nominal: list[tuple[str, str]] = []
    numeric: list[tuple[float, float]] = []
    for row in rows:
        left, right = row.get(human), row.get(judge)
        if is_missing(left, missing) or is_missing(right, missing):
            continue
        nominal.append((text(left), text(right)))
        left_num, right_num = as_number(left), as_number(right)
        if left_num is not None and right_num is not None:
            numeric.append((left_num, right_num))
    exact = sum(left == right for left, right in nominal)
    result: dict[str, Any] = {
        "paired_n": len(nominal),
        "exact_agreement": rounded(exact / len(nominal)) if nominal else None,
        "cohens_kappa_nominal": nominal_kappa(nominal),
    }
    if numeric:
        differences = [right - left for left, right in numeric]
        result["numeric"] = {
            "paired_n": len(numeric),
            "mean_judge_minus_human": rounded(statistics.fmean(differences)),
            "mean_absolute_error": rounded(statistics.fmean(abs(value) for value in differences)),
            "pearson_correlation": pearson(numeric),
        }
    return result


def mcnemar_exact(regressions: int, improvements: int) -> float | None:
    discordant = regressions + improvements
    if discordant == 0:
        return 1.0
    lower = min(regressions, improvements)
    probability = sum(math.comb(discordant, k) for k in range(lower + 1)) / (2**discordant)
    return rounded(min(1.0, 2 * probability))


def comparison_profile(
    rows: list[dict[str, Any]],
    version_column: str,
    control: str,
    experiment: str,
    id_column: str | None,
    pass_column: str,
    classify,
    missing: set[str],
) -> dict[str, Any]:
    control_rows = [row for row in rows if text(row.get(version_column)) == control]
    experiment_rows = [row for row in rows if text(row.get(version_column)) == experiment]
    control_stats = pass_profile(control_rows, pass_column, classify, missing)
    experiment_stats = pass_profile(experiment_rows, pass_column, classify, missing)
    left, right = control_stats["pass_rate"], experiment_stats["pass_rate"]
    result: dict[str, Any] = {
        "control": {"value": control, **control_stats},
        "experiment": {"value": experiment, **experiment_stats},
        "absolute_change_percentage_points": rounded((right - left) * 100)
        if left is not None and right is not None
        else None,
        "relative_change": rounded((right - left) / left)
        if left not in (None, 0) and right is not None
        else None,
    }
    if not id_column:
        return result
    index: dict[str, dict[str, list[bool]]] = defaultdict(lambda: defaultdict(list))
    for row in control_rows + experiment_rows:
        identifier = row.get(id_column)
        outcome = classify(row.get(pass_column))
        if is_missing(identifier, missing) or outcome is None:
            continue
        index[text(identifier)][text(row.get(version_column))].append(outcome)
    pairs = []
    duplicate_pairs = 0
    for versions in index.values():
        if len(versions.get(control, [])) == 1 and len(versions.get(experiment, [])) == 1:
            pairs.append((versions[control][0], versions[experiment][0]))
        elif control in versions and experiment in versions:
            duplicate_pairs += 1
    regressions = sum(left and not right for left, right in pairs)
    improvements = sum(not left and right for left, right in pairs)
    result["paired"] = {
        "paired_n": len(pairs),
        "improvements_n": improvements,
        "regressions_n": regressions,
        "unchanged_n": len(pairs) - improvements - regressions,
        "duplicate_pair_ids_excluded_n": duplicate_pairs,
        "paired_change_percentage_points": rounded((improvements - regressions) / len(pairs) * 100)
        if pairs
        else None,
        "mcnemar_exact_two_sided_p": mcnemar_exact(regressions, improvements),
    }
    return result


def parse_values(raw: str | None) -> set[str]:
    if not raw:
        return set()
    return {item.strip().casefold() for item in raw.split(",") if item.strip()}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="只读分析 CSV/TSV/JSON/JSONL/XLSX 测评结果；字段口径必须由调用者显式确认。"
    )
    parser.add_argument("input", help="输入文件")
    parser.add_argument("--sheet", help="XLSX 工作表；不填时读取活动工作表")
    parser.add_argument("--id-column", help="测试样本唯一 ID 字段")
    parser.add_argument("--pass-column", help="通过/失败状态字段")
    parser.add_argument("--pass-values", help="逗号分隔的通过值；使用 --pass-column 时必填")
    parser.add_argument("--fail-values", help="逗号分隔的失败值；使用 --pass-column 时必填")
    parser.add_argument("--score-column", help="总体数值分数字段")
    parser.add_argument("--human-score-column", help="人工评分或标签字段")
    parser.add_argument("--judge-score-column", help="Judge 评分或标签字段")
    parser.add_argument("--version-column", help="模型或系统版本字段")
    parser.add_argument("--control-value", help="对照组版本值")
    parser.add_argument("--experiment-value", help="实验组版本值")
    parser.add_argument("--group-by", action="append", default=[], help="分组字段，可重复")
    parser.add_argument("--missing-value", action="append", default=[], help="额外缺失值，可重复")
    parser.add_argument("--output", help="单独输出 JSON 文件；默认写 stdout")
    parser.add_argument("--overwrite", action="store_true", help="允许覆盖已有的输出文件，不影响输入")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    path = Path(args.input).expanduser().resolve()
    if not path.is_file():
        parser.error(f"输入文件不存在：{path}")
    if args.pass_column and (not args.pass_values or not args.fail_values):
        parser.error("使用 --pass-column 时必须同时提供 --pass-values 和 --fail-values。")
    if bool(args.control_value) != bool(args.experiment_value):
        parser.error("--control-value 与 --experiment-value 必须同时提供。")
    if (args.control_value or args.experiment_value) and not (args.version_column and args.pass_column):
        parser.error("版本比较还需要 --version-column 和 --pass-column。")
    if bool(args.human_score_column) != bool(args.judge_score_column):
        parser.error("--human-score-column 与 --judge-score-column 必须同时提供。")

    try:
        rows, source = read_rows(path, args.sheet)
        columns = ordered_columns(rows)
        configured = [
            args.id_column,
            args.pass_column,
            args.score_column,
            args.human_score_column,
            args.judge_score_column,
            args.version_column,
            *args.group_by,
        ]
        missing_columns = sorted({name for name in configured if name and name not in columns})
        if missing_columns:
            raise UserError(f"指定字段不存在：{missing_columns}")

        missing_values = {value.casefold() for value in args.missing_value}
        pass_values = parse_values(args.pass_values)
        fail_values = parse_values(args.fail_values)
        overlap = pass_values & fail_values
        if overlap:
            raise UserError(f"通过值和失败值重叠：{sorted(overlap)}")
        classify = make_classifier(pass_values, fail_values, missing_values)

        missing_by_column = {}
        for column in columns:
            count = sum(is_missing(row.get(column), missing_values) for row in rows)
            missing_by_column[column] = {
                "missing_n": count,
                "missing_rate": rounded(count / len(rows)) if rows else None,
            }

        canonical = [json.dumps(row, sort_keys=True, ensure_ascii=False, default=str) for row in rows]
        exact_duplicates = len(canonical) - len(set(canonical))
        duplicate_key_records = None
        if args.id_column:
            keys = []
            for row in rows:
                identifier = row.get(args.id_column)
                if is_missing(identifier, missing_values):
                    continue
                key = (text(identifier), text(row.get(args.version_column))) if args.version_column else (text(identifier),)
                keys.append(key)
            counts = Counter(keys)
            duplicate_key_records = sum(count - 1 for count in counts.values() if count > 1)

        report: dict[str, Any] = {
            "source": {"input_name": path.name, **source},
            "scope": {
                "row_count": len(rows),
                "column_count": len(columns),
                "columns": columns,
                "unique_id_count": len(
                    {
                        text(row.get(args.id_column))
                        for row in rows
                        if args.id_column and not is_missing(row.get(args.id_column), missing_values)
                    }
                )
                if args.id_column
                else None,
            },
            "quality": {
                "exact_duplicate_rows_n": exact_duplicates,
                "duplicate_key_records_n": duplicate_key_records,
                "missing_by_column": missing_by_column,
            },
        }

        if args.pass_column:
            report["outcomes"] = {"overall": pass_profile(rows, args.pass_column, classify, missing_values)}
        numeric_columns = [
            column
            for column in (args.score_column, args.human_score_column, args.judge_score_column)
            if column
        ]
        if numeric_columns:
            report["numeric_profiles"] = {
                column: numeric_profile(rows, column, missing_values) for column in dict.fromkeys(numeric_columns)
            }

        group_columns = list(dict.fromkeys([*args.group_by, args.version_column] if args.version_column else args.group_by))
        if group_columns:
            report["groups"] = {}
            for column in group_columns:
                entries = {}
                for label, subset in sorted(group_rows(rows, column, missing_values).items()):
                    entry: dict[str, Any] = {"row_count": len(subset)}
                    if args.pass_column:
                        entry["outcomes"] = pass_profile(subset, args.pass_column, classify, missing_values)
                    if args.score_column:
                        entry["score"] = numeric_profile(subset, args.score_column, missing_values)
                    entries[label] = entry
                report["groups"][column] = entries

        if args.human_score_column and args.judge_score_column:
            report["human_judge_agreement"] = agreement_profile(
                rows, args.human_score_column, args.judge_score_column, missing_values
            )

        if args.control_value and args.experiment_value:
            report["version_comparison"] = comparison_profile(
                rows,
                args.version_column,
                args.control_value,
                args.experiment_value,
                args.id_column,
                args.pass_column,
                classify,
                missing_values,
            )

        warnings = []
        if len(rows) < 30:
            warnings.append("总记录数少于 30；比例和版本差异的不确定性可能很大。")
        if args.pass_column and report["outcomes"]["overall"]["unrecognized_n"]:
            warnings.append("存在未识别状态值；这些记录未进入通过率分母。")
        if args.control_value and args.experiment_value and not args.id_column:
            warnings.append("未提供 ID 字段；版本比较无法进行配对分析。")
        report["warnings"] = warnings

        rendered = json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False)
        if args.output:
            output = Path(args.output).expanduser().resolve()
            if output == path:
                raise UserError("输出文件不能与输入文件相同。")
            if output.exists() and not args.overwrite:
                raise UserError("输出文件已存在；更换路径或显式使用 --overwrite。")
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(rendered + "\n", encoding="utf-8")
            print(f"分析结果已写入：{output}")
        else:
            print(rendered)
        return 0
    except UserError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
